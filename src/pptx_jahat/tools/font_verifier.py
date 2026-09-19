"""
font_verifier.py
================
Verifies template typography against installed system fonts and project fonts.
Provides intelligent fallback recommendations, user fallback resolution, and
deep OpenXML / DrawingML font replacement for PowerPoint presentations.
"""

from __future__ import annotations

import os
import sys
import re
import logging
from pathlib import Path
from typing import Dict, List, Set, Tuple, Optional, Any
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

try:
    from fontTools.ttLib import TTFont, TTCollection
    HAS_FONTTOOLS = True
except Exception:
    HAS_FONTTOOLS = False

log = logging.getLogger("pptx_jahat.font_verifier")

# Common keywords for Persian / Arabic script typefaces
PERSIAN_ARABIC_KEYWORDS = [
    "yekan", "nazanin", "vazir", "vazirmatn", "amuzeh", "titr", "shabnam", "sahel",
    "mitra", "zar", "roya", "lotus", "fanum", "arabic", "farsi", "persian",
    "parastoo", "samim", "tanha", "shiraz", "esfehan", "morvarid",
    "koodak", "homa", "khoram", "sina", "tabassom", "besmellah", "khodkar",
    "suls", "thuluth", "nastaliq", "shekasteh", "naskh", "diwani", "kufi",
    "kufic", "riqah", "roqaa", "aldhabi", "amiri", "cairo", "tajawal",
    "harmattan", "lateef", "scheherazade", "markazi", "katibeh", "aref",
    "anjoman", "alef", "dana", "peyda", "kalameh", "doran", "morabba",
    "estedad", "rokh", "gandom", "rastik", "dima", "arshia", "aseman",
    "badr", "baran", "bardiya", "compset", "davat", "elham", "farnaz",
    "ferdos", "hamid", "helal", "jaded", "jalal", "jamil", "kamran",
    "karim", "kaveh", "kayhan", "mahsa", "majid", "makhsoos", "marooff",
    "masjed", "medad", "mehr", "narm", "nasim", "nikoo", "sarah",
    "sepideh", "setareh", "shadi", "shafigh", "siamak", "sooreh", "soufi",
    "tavallod", "tehran", "vahid", "yagut", "yas", "yashar", "ziba",
    "casablanca", "adobe arabic", "traditional arabic", "simplified arabic",
    "urdu", "pashto", "kurdish", "sansx"
]

# Standard preferred Persian fallback fonts in priority order
PREFERRED_PERSIAN_FALLBACKS = [
    "IRANYekanXFaNum",
    "IRANYekanXFaNum Heavy",
    "IRANYekanXFaNum ExtraBlack",
    "IRANYekanX ExtraBlack",
    "Vazirmatn",
    "B Nazanin",
    "Tahoma",
    "Segoe UI",
    "Arial"
]

# Standard preferred Latin / Universal fallback fonts in priority order
PREFERRED_LATIN_FALLBACKS = [
    "Segoe UI",
    "Calibri",
    "Arial",
    "Montserrat",
    "Helvetica",
    "Tahoma"
]

_CACHED_SYSTEM_FONTS: Optional[Dict[str, Any]] = None


def normalize_font_name(name: str) -> str:
    """Normalizes font name for fuzzy comparison by removing spaces, hyphens, distributor tags, and casing."""
    if not name:
        return ""
    # Remove file extensions if present
    cleaned = re.sub(r'\.(ttf|otf|ttc)$', '', name, flags=re.IGNORECASE)
    # Remove distributor / pack tags like [@mimvid], (MRT), [fontyab]
    cleaned = re.sub(r'\[.*?\]|\(.*?\)|_mrt_|_mrt|mrt_', '', cleaned, flags=re.IGNORECASE)
    # Remove common style suffixes for family comparison
    cleaned = re.sub(r'[\s\-_]+(bold|italic|regular|light|medium|semibold|demibold|heavy|black|extrablack|thin|oblique|ultralight|extralight|narrow|condensed|book|roman)', '', cleaned, flags=re.IGNORECASE)
    # Remove Persian/Arabic numeral / font format suffixes (FaNum, FD, WOL, RD, etc.)
    cleaned = re.sub(r'[\s\-_]*(fanum|vfanum|fd|wol|rd)$', '', cleaned, flags=re.IGNORECASE)
    # Normalize variable font markers (XV, VF -> X)
    cleaned = re.sub(r'(xv|vf)$', 'x', cleaned, flags=re.IGNORECASE)
    return re.sub(r'[^a-zA-Z0-9\u0600-\u06FF]', '', cleaned).lower()


def is_persian_or_arabic_font(font_name: str) -> bool:
    """Determines whether a font name is intended for Persian or Arabic scripts."""
    if not font_name:
        return False
    lower = font_name.lower().strip()
    # Explicit exclusions for common Latin typefaces
    if any(latin in lower for latin in (
        "arial", "calibri", "segoe", "times", "helvetica", "montserrat",
        "courier", "consolas", "verdana", "roboto", "georgia", "trebuchet",
        "comic", "impact", "tahoma", "cambria", "garamond", "century gothic",
        "lucida", "franklin", "palatino", "bookman", "candara", "corbel",
        "constantia", "optima", "futura", "baskerville", "caslon", "bodoni",
        "didot", "gill sans", "aptos", "inter", "geist"
    )):
        return False
    # Check for actual Arabic/Persian Unicode characters in name
    for ch in font_name:
        if ('\u0600' <= ch <= '\u06FF' or '\u0750' <= ch <= '\u077F'
                or '\uFB50' <= ch <= '\uFDFF' or '\uFE70' <= ch <= '\uFEFF'):
            return True
    if (
        lower.startswith("b ") or lower.startswith("b-") or lower.startswith("b_") or
        lower.startswith("ir ") or lower.startswith("ir_") or lower.startswith("ir-") or lower.startswith("iran") or
        lower.startswith("2 ") or lower.startswith("2-") or lower.startswith("2_") or
        lower.startswith("mrt") or lower.startswith("_mrt") or
        lower.startswith("mj ") or lower.startswith("mj_") or
        lower.startswith("a ") or lower.startswith("a_") or
        lower.startswith("w_")
    ):
        return True
    return any(k in lower for k in PERSIAN_ARABIC_KEYWORDS)


def _scan_windows_gdi() -> Tuple[Set[str], Set[str]]:
    """Uses Windows GDI EnumFontFamiliesExW to query all registered fonts (<5ms)."""
    families: Set[str] = set()
    full_names: Set[str] = set()
    if os.name != "nt":
        return families, full_names

    try:
        import ctypes
        from ctypes import wintypes

        user32 = ctypes.windll.user32
        gdi32 = ctypes.windll.gdi32

        LF_FACESIZE = 32
        LF_FULLFACESIZE = 64

        class LOGFONTW(ctypes.Structure):
            _fields_ = [
                ('lfHeight', wintypes.LONG), ('lfWidth', wintypes.LONG),
                ('lfEscapement', wintypes.LONG), ('lfOrientation', wintypes.LONG),
                ('lfWeight', wintypes.LONG), ('lfItalic', wintypes.BYTE),
                ('lfUnderline', wintypes.BYTE), ('lfStrikeOut', wintypes.BYTE),
                ('lfCharSet', wintypes.BYTE), ('lfOutPrecision', wintypes.BYTE),
                ('lfClipPrecision', wintypes.BYTE), ('lfQuality', wintypes.BYTE),
                ('lfPitchAndFamily', wintypes.BYTE), ('lfFaceName', wintypes.WCHAR * LF_FACESIZE)
            ]

        class ENUMLOGFONTEXW(ctypes.Structure):
            _fields_ = [
                ('elfLogFont', LOGFONTW),
                ('elfFullName', wintypes.WCHAR * LF_FULLFACESIZE),
                ('elfStyle', wintypes.WCHAR * LF_FACESIZE),
                ('elfScript', wintypes.WCHAR * LF_FACESIZE)
            ]

        FONTENUMPROC = ctypes.WINFUNCTYPE(
            ctypes.c_int,
            ctypes.POINTER(ENUMLOGFONTEXW),
            ctypes.c_void_p,
            wintypes.DWORD,
            wintypes.LPARAM
        )

        def enum_callback(lpelfe, lpntme, FontType, lParam):
            face = lpelfe.contents.elfLogFont.lfFaceName
            full = lpelfe.contents.elfFullName
            if face:
                families.add(face.strip())
            if full:
                full_names.add(full.strip())
            return 1

        cb = FONTENUMPROC(enum_callback)
        hdc = user32.GetDC(None)
        if hdc:
            try:
                lf = LOGFONTW()
                lf.lfCharSet = 1  # DEFAULT_CHARSET
                gdi32.EnumFontFamiliesExW(hdc, ctypes.byref(lf), cb, 0, 0)
            finally:
                user32.ReleaseDC(None, hdc)
    except Exception as ex:
        log.debug("GDI EnumFontFamiliesExW scan failed or unavailable: %s", ex)

    return families, full_names


def _scan_windows_registry() -> Tuple[Set[str], Dict[str, str]]:
    """
    Reads Windows Registry for machine-wide and per-user fonts:
    HKLM and HKCU Software\\Microsoft\\Windows NT\\CurrentVersion\\Fonts.
    """
    font_names: Set[str] = set()
    file_map: Dict[str, str] = {}
    if os.name != "nt":
        return font_names, file_map

    try:
        import winreg
        windir_fonts = Path(os.environ.get("WINDIR", "C:\\Windows")) / "Fonts"

        for root in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
            for subkey in (
                r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts",
                r"SOFTWARE\WOW6432Node\Microsoft\Windows NT\CurrentVersion\Fonts"
            ):
                try:
                    key = winreg.OpenKey(root, subkey)
                    for i in range(winreg.QueryInfoKey(key)[1]):
                        try:
                            vname, val, _ = winreg.EnumValue(key, i)
                            if not vname:
                                continue
                            # Clean (TrueType), (OpenType), etc.
                            cleaned = re.sub(r'\s*\((TrueType|OpenType|PostScript|All res|\d+)\)$', '', vname, flags=re.IGNORECASE).strip()
                            # Resolve file path
                            val_str = str(val).strip()
                            resolved_path = val_str
                            if val_str and not Path(val_str).is_absolute() and (windir_fonts / val_str).exists():
                                resolved_path = str((windir_fonts / val_str).resolve())

                            # Split names joined with '&' (e.g. "Cambria & Cambria Math")
                            for part in re.split(r'\s*&\s*', cleaned):
                                p_clean = part.strip()
                                if p_clean:
                                    font_names.add(p_clean)
                                    if resolved_path and Path(resolved_path).exists():
                                        file_map[p_clean.lower()] = resolved_path
                        except Exception:
                            pass
                    winreg.CloseKey(key)
                except Exception:
                    pass
    except Exception as ex:
        log.debug("Windows registry font scan skipped: %s", ex)

    return font_names, file_map


def _scan_font_files_metadata(directories: List[Path]) -> Tuple[Set[str], Set[str], Dict[str, str]]:
    """
    Scans TTF, OTF, and TTC font files in directories using fontTools to extract
    true Family, Full Name, Typographic Family (NameID 16), and WWS Family (NameID 21).
    """
    families: Set[str] = set()
    full_names: Set[str] = set()
    file_map: Dict[str, str] = {}

    if not HAS_FONTTOOLS:
        return families, full_names, file_map

    for d in directories:
        if not d.exists():
            continue

        # Single font files (.ttf, .otf)
        font_files = list(d.glob("*.ttf")) + list(d.glob("*.otf")) + list(d.glob("*.TTF")) + list(d.glob("*.OTF"))
        for f in font_files:
            file_path_str = str(f.resolve())
            try:
                tt = TTFont(str(f), fontNumber=0, lazy=True)
                nt = tt.get("name")
                if nt is not None:
                    for r in nt.names:
                        try:
                            val = r.toUnicode().strip()
                            if not val or len(val) < 2:
                                continue
                            if r.nameID in (1, 16, 21):
                                families.add(val)
                                file_map[val.lower()] = file_path_str
                            elif r.nameID in (4, 6):
                                full_names.add(val)
                                file_map[val.lower()] = file_path_str
                        except Exception:
                            pass
            except Exception:
                # Fallback to filename stem
                stem = f.stem.replace("-", " ").replace("_", " ").strip()
                if stem:
                    families.add(stem)
                    file_map[stem.lower()] = file_path_str

        # TrueType Collections (.ttc)
        ttc_files = list(d.glob("*.ttc")) + list(d.glob("*.TTC"))
        for f in ttc_files:
            file_path_str = str(f.resolve())
            try:
                coll = TTCollection(str(f))
                for font in coll:
                    nt = font.get("name")
                    if nt is not None:
                        for r in nt.names:
                            try:
                                val = r.toUnicode().strip()
                                if not val or len(val) < 2:
                                    continue
                                if r.nameID in (1, 16, 21):
                                    families.add(val)
                                    file_map[val.lower()] = file_path_str
                                elif r.nameID in (4, 6):
                                    full_names.add(val)
                                    file_map[val.lower()] = file_path_str
                            except Exception:
                                pass
            except Exception:
                stem = f.stem.replace("-", " ").replace("_", " ").strip()
                if stem:
                    families.add(stem)
                    file_map[stem.lower()] = file_path_str

    return families, full_names, file_map


def get_available_system_fonts(force_rescan: bool = False) -> Dict[str, Any]:
    """
    Comprehensive multi-tier font discovery engine:
    1. Windows GDI subsystem (EnumFontFamiliesExW)
    2. Windows Registry (HKLM, HKCU, WOW6432Node)
    3. User fonts (%LOCALAPPDATA%\\Microsoft\\Windows\\Fonts)
    4. System fonts (%WINDIR%\\Fonts, /usr/share/fonts, ~/Library/Fonts, etc.)
    5. Project fonts (assets/fonts)
    6. OpenType metadata parsing (Typographic Family, Full Name, WWS Family)
    7. Base family derivation & distributor tag normalizations
    8. Matplotlib fallback
    """
    global _CACHED_SYSTEM_FONTS
    if _CACHED_SYSTEM_FONTS is not None and not force_rescan:
        return _CACHED_SYSTEM_FONTS

    found_families: Set[str] = set()
    exact_names: Set[str] = set()
    normalized_map: Dict[str, str] = {}
    font_paths: Dict[str, str] = {}

    def register_name(raw_name: str, file_path: Optional[str] = None) -> None:
        if not raw_name:
            return
        name = raw_name.strip()
        if not name or len(name) < 2:
            return
        # If passed raw filename ending in font extension, strip it
        if name.lower().endswith(('.ttf', '.otf', '.ttc')):
            name = Path(name).stem

        found_families.add(name)
        exact_names.add(name.lower())
        norm = normalize_font_name(name)
        if norm and norm not in normalized_map:
            normalized_map[norm] = name
        if file_path and name.lower() not in font_paths:
            font_paths[name.lower()] = file_path

        # 1. Strip distributor / channel tags: e.g. "IRANYekan  [ @mimvid ] Bold" -> "IRANYekan Bold"
        dist_cleaned = re.sub(r'\s*(\[.*?\]|\(.*?\)|_mrt_|_mrt|mrt_)\s*', ' ', name, flags=re.IGNORECASE).strip()
        if dist_cleaned and dist_cleaned != name and len(dist_cleaned) >= 2:
            found_families.add(dist_cleaned)
            exact_names.add(dist_cleaned.lower())
            norm_dist = normalize_font_name(dist_cleaned)
            if norm_dist and norm_dist not in normalized_map:
                normalized_map[norm_dist] = dist_cleaned
            if file_path and dist_cleaned.lower() not in font_paths:
                font_paths[dist_cleaned.lower()] = file_path

        # 2. Strip common weight / style suffixes to register the base family:
        # e.g. "Montserrat Black" -> "Montserrat", "IRANYekanXFaNum ExtraBlack" -> "IRANYekanXFaNum"
        base_name = re.sub(
            r'[\s\-_]+(bold|italic|regular|light|medium|semibold|demibold|heavy|black|extrablack|thin|oblique|ultralight|extralight|narrow|condensed|book|roman)$',
            '', name, flags=re.IGNORECASE
        ).strip()
        if base_name and base_name != name and len(base_name) >= 2:
            found_families.add(base_name)
            exact_names.add(base_name.lower())
            norm_base = normalize_font_name(base_name)
            if norm_base and norm_base not in normalized_map:
                normalized_map[norm_base] = base_name
            if file_path and base_name.lower() not in font_paths:
                font_paths[base_name.lower()] = file_path

        # 3. Base name from distributor cleaned
        if dist_cleaned and dist_cleaned != name:
            base_dist = re.sub(
                r'[\s\-_]+(bold|italic|regular|light|medium|semibold|demibold|heavy|black|extrablack|thin|oblique|ultralight|extralight|narrow|condensed|book|roman)$',
                '', dist_cleaned, flags=re.IGNORECASE
            ).strip()
            if base_dist and base_dist != dist_cleaned and len(base_dist) >= 2:
                found_families.add(base_dist)
                exact_names.add(base_dist.lower())
                norm_bd = normalize_font_name(base_dist)
                if norm_bd and norm_bd not in normalized_map:
                    normalized_map[norm_bd] = base_dist
                if file_path and base_dist.lower() not in font_paths:
                    font_paths[base_dist.lower()] = file_path

    # Tier 1: Windows GDI subsystem scan
    gdi_families, gdi_full = _scan_windows_gdi()
    for fam in gdi_families:
        register_name(fam)
    for full in gdi_full:
        register_name(full)

    # Tier 2: Windows Registry scan
    reg_names, reg_files = _scan_windows_registry()
    for rname in reg_names:
        register_name(rname, reg_files.get(rname.lower()))

    # Tier 3: Font directories (Windows system, user AppData, Linux, macOS, project)
    font_dirs: List[Path] = []
    if os.name == "nt":
        windir = os.environ.get("WINDIR", "C:\\Windows")
        font_dirs.append(Path(windir) / "Fonts")
        localapp = os.environ.get("LOCALAPPDATA", "")
        if localapp:
            font_dirs.append(Path(localapp) / "Microsoft" / "Windows" / "Fonts")
        appdata = os.environ.get("APPDATA", "")
        if appdata:
            font_dirs.append(Path(appdata) / "Microsoft" / "Windows" / "Fonts")
    else:
        font_dirs.extend([
            Path("/usr/share/fonts"),
            Path("/usr/local/share/fonts"),
            Path("~/.fonts").expanduser(),
            Path("~/.local/share/fonts").expanduser(),
            Path("/System/Library/Fonts"),
            Path("/Library/Fonts"),
            Path("~/Library/Fonts").expanduser(),
        ])

    # Project assets/fonts
    project_fonts_dir = Path(__file__).resolve().parent.parent.parent.parent / "assets" / "fonts"
    if project_fonts_dir.exists():
        font_dirs.append(project_fonts_dir)

    # Parse metadata from font files via fontTools
    dir_families, dir_full, dir_files = _scan_font_files_metadata(font_dirs)
    for fam in dir_families:
        register_name(fam, dir_files.get(fam.lower()))
    for full in dir_full:
        register_name(full, dir_files.get(full.lower()))

    # Tier 4: Matplotlib fontManager fallback
    try:
        import matplotlib.font_manager as fm
        for entry in getattr(fm.fontManager, "ttflist", []):
            if entry.name:
                register_name(entry.name, getattr(entry, "fname", None))
    except Exception as ex:
        log.debug("matplotlib font_manager fallback scan skipped: %s", ex)

    # Filter out empty or pure-symbol names
    valid_families = [
        f for f in found_families
        if f and len(f) >= 2 and not f.startswith((".", "-", "+", "$", "@", "_"))
    ]
    sorted_families = sorted(list(set(valid_families)), key=lambda s: s.lower())

    # Partition into Persian/Arabic vs Latin
    persian_fonts = [f for f in sorted_families if is_persian_or_arabic_font(f)]
    latin_fonts = [f for f in sorted_families if not is_persian_or_arabic_font(f)]

    _CACHED_SYSTEM_FONTS = {
        "families": sorted_families,
        "exact_names": exact_names,
        "normalized_map": normalized_map,
        "persian_fonts": persian_fonts,
        "latin_fonts": latin_fonts,
        "font_paths": font_paths,
        "total_count": len(sorted_families)
    }
    log.info(
        "Discovered %d system fonts (%d Persian/Arabic, %d Latin/Universal).",
        len(sorted_families), len(persian_fonts), len(latin_fonts)
    )
    return _CACHED_SYSTEM_FONTS


def is_font_available(font_name: str, available_fonts: Optional[Dict[str, Any]] = None) -> bool:
    """
    Checks if a given font name is installed and available on the system.
    Office theme placeholders like '+mj-lt' or '+mn-lt' return True.
    """
    if not font_name:
        return True
    
    clean_name = font_name.strip()
    if not clean_name:
        return True
    
    # Office theme placeholders (e.g. +mj-lt, +mn-lt, +mj-ea, +mn-ea, +mj-cs, +mn-cs)
    if clean_name.startswith("+"):
        return True

    catalog = available_fonts or get_available_system_fonts()
    exact_set: Set[str] = catalog.get("exact_names", set())
    families: List[str] = catalog.get("families", [])
    norm_map: Dict[str, str] = catalog.get("normalized_map", {})

    lower_name = clean_name.lower()

    # 1. Exact match (case-insensitive)
    if lower_name in exact_set:
        return True

    # 2. Case-insensitive match in families
    for f in families:
        if f.lower() == lower_name:
            return True

    # 3. Normalized match (ignoring spaces, hyphens, distributor tags, and common weights)
    norm = normalize_font_name(clean_name)
    if norm and norm in norm_map:
        return True

    # 4. Clean distributor tags and test
    clean_dist = re.sub(r'\[.*?\]|\(.*?\)|_mrt_|_mrt|mrt_', '', clean_name, flags=re.IGNORECASE).strip()
    if clean_dist and clean_dist.lower() in exact_set:
        return True
    norm_dist = normalize_font_name(clean_dist)
    if norm_dist and norm_dist in norm_map:
        return True

    # 5. Base family match without weight/style
    clean_no_weight = re.sub(
        r'[\s\-_]+(bold|italic|regular|heavy|black|extrablack|light|medium|semibold|demibold|vfanum|fanum|thin|oblique|narrow|condensed)$',
        '', lower_name, flags=re.IGNORECASE
    ).strip()
    if clean_no_weight:
        if clean_no_weight in exact_set:
            return True
        norm_nw = normalize_font_name(clean_no_weight)
        if norm_nw and norm_nw in norm_map:
            return True

    # 6. Word-boundary family prefix match
    if len(clean_no_weight) >= 3:
        for f in families:
            f_lower = f.lower()
            if f_lower == clean_no_weight:
                return True
            # e.g. presentation has "Montserrat", installed is "Montserrat Black"
            # or presentation has "IRANSansXFaNum DemiBold", installed is "IRANSansX"
            if f_lower.startswith(clean_no_weight + " ") or clean_no_weight.startswith(f_lower + " "):
                return True

    return False
    if not clean_name:
        return True
    
    # Office theme placeholders (e.g. +mj-lt, +mn-lt, +mj-ea, +mn-ea, +mj-cs)
    if clean_name.startswith("+"):
        return True

    catalog = available_fonts or get_available_system_fonts()
    exact_set: Set[str] = catalog.get("exact_names", set())
    families: List[str] = catalog.get("families", [])
    norm_map: Dict[str, str] = catalog.get("normalized_map", {})

    lower_name = clean_name.lower()

    # 1. Exact match (case-insensitive)
    if lower_name in exact_set:
        return True

    # 2. Case-insensitive match in families
    for f in families:
        if f.lower() == lower_name:
            return True

    # 3. Normalized match (ignoring spaces, hyphens, and common weights)
    norm = normalize_font_name(clean_name)
    if norm in norm_map:
        return True

    # 4. Partial family prefix match
    clean_no_weight = re.sub(r'[\s\-_]+(bold|italic|regular|heavy|black|extrablack|light|medium|semibold|demibold|vfanum|fanum)$', '', lower_name, flags=re.IGNORECASE)
    for f in families:
        f_lower = f.lower()
        if f_lower.startswith(clean_no_weight) or clean_no_weight.startswith(f_lower):
            return True

    return False


def recommend_font_fallback(font_name: str, available_fonts: Optional[Dict[str, Any]] = None) -> str:
    """
    Recommends the best available system font as a fallback for a missing font.
    """
    catalog = available_fonts or get_available_system_fonts()
    families_lower = {f.lower(): f for f in catalog.get("families", [])}
    persian_available = catalog.get("persian_fonts", [])
    latin_available = catalog.get("latin_fonts", [])

    is_persian = is_persian_or_arabic_font(font_name)
    name_lower = font_name.lower()
    is_heavy = any(w in name_lower for w in ("heavy", "black", "extrablack", "bold"))

    if is_persian:
        # Check heavy variants first if original is heavy
        if is_heavy:
            for heavy_cand in ("IRANYekanXFaNum Heavy", "IRANYekanXFaNum ExtraBlack", "IRANYekanX ExtraBlack", "B Titr", "B Nazanin Bold"):
                if heavy_cand.lower() in families_lower:
                    return families_lower[heavy_cand.lower()]

        # Try standard preferred Persian fallbacks
        for cand in PREFERRED_PERSIAN_FALLBACKS:
            if cand.lower() in families_lower:
                return families_lower[cand.lower()]

        # Any installed Persian font
        if persian_available:
            return persian_available[0]

        # Universal fallback
        return "Tahoma"
    else:
        # Latin / Universal font
        for cand in PREFERRED_LATIN_FALLBACKS:
            if cand.lower() in families_lower:
                return families_lower[cand.lower()]

        if latin_available:
            return latin_available[0]

        return "Segoe UI"


def verify_fonts_inventory(fonts: List[str]) -> Dict[str, Any]:
    """
    Verifies a list of fonts against the system catalog.
    Returns installed fonts, missing fonts, recommendations, and available system fonts.
    """
    catalog = get_available_system_fonts()
    
    unique_fonts = sorted(list(set(f.strip() for f in fonts if f and f.strip())))
    installed: List[str] = []
    missing: List[str] = []
    recommendations: Dict[str, str] = {}

    for f in unique_fonts:
        # Filter theme placeholders
        if f.startswith("+"):
            continue
        if is_font_available(f, catalog):
            installed.append(f)
        else:
            missing.append(f)
            recommendations[f] = recommend_font_fallback(f, catalog)

    all_installed = len(missing) == 0

    return {
        "all_installed": all_installed,
        "total_checked": len(unique_fonts),
        "installed_fonts": installed,
        "missing_fonts": missing,
        "recommendations": recommendations,
        "available_persian": catalog.get("persian_fonts", []),
        "available_latin": catalog.get("latin_fonts", []),
        "system_fonts": catalog.get("families", [])
    }


def apply_font_fallbacks_to_presentation(
    prs: Any,
    font_fallbacks: Dict[str, str]
) -> int:
    """
    Walks through all slides, shapes, text frames, paragraphs, runs, and tables
    in a Presentation and replaces any font present in font_fallbacks with its target fallback.
    Returns the count of modified runs and elements.
    """
    if not font_fallbacks:
        return 0

    # Build case-insensitive and normalized mapping
    mapping: Dict[str, str] = {}
    normalized_mapping: Dict[str, str] = {}
    for src, dst in font_fallbacks.items():
        if src and dst and src.strip() != dst.strip():
            src_clean = src.strip()
            dst_clean = dst.strip()
            mapping[src_clean.lower()] = dst_clean
            normalized_mapping[normalize_font_name(src_clean)] = dst_clean

    if not mapping:
        return 0

    modified_count = 0

    def resolve_target(font_face: Optional[str]) -> Optional[str]:
        if not font_face:
            return None
        lower = font_face.strip().lower()
        if lower in mapping:
            return mapping[lower]
        norm = normalize_font_name(font_face)
        if norm in normalized_mapping:
            return normalized_mapping[norm]
        return None

    def process_text_frame(tf: Any) -> None:
        nonlocal modified_count
        if not tf:
            return
        for p in tf.paragraphs:
            # Check paragraph default run properties (defRPr)
            if hasattr(p, "_p"):
                pPr = p._p.find("{http://schemas.openxmlformats.org/drawingml/2006/main}pPr")
                if pPr is not None:
                    defRPr = pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr")
                    if defRPr is not None:
                        for tag in ("cs", "latin", "ea"):
                            elem = defRPr.find(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag}")
                            if elem is not None and elem.get("typeface"):
                                target = resolve_target(elem.get("typeface"))
                                if target:
                                    elem.set("typeface", target)
                                    modified_count += 1

            # Check individual runs
            for r in p.runs:
                # python-pptx font.name
                if r.font and r.font.name:
                    target = resolve_target(r.font.name)
                    if target:
                        r.font.name = target
                        modified_count += 1

                # DrawingML rPr typeface (cs, latin, ea)
                if hasattr(r, "_r"):
                    rPr = r._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                    if rPr is not None:
                        for tag in ("cs", "latin", "ea"):
                            elem = rPr.find(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag}")
                            if elem is not None and elem.get("typeface"):
                                target = resolve_target(elem.get("typeface"))
                                if target:
                                    elem.set("typeface", target)
                                    modified_count += 1

    def process_shape(shape: Any) -> None:
        if shape.has_text_frame:
            process_text_frame(shape.text_frame)

        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if cell.text_frame:
                        process_text_frame(cell.text_frame)

        # Handle grouped shapes
        if hasattr(shape, "shapes"):
            for child in shape.shapes:
                process_shape(child)

    for slide in prs.slides:
        for shape in slide.shapes:
            process_shape(shape)

    log.info("Applied font fallbacks to presentation: %d font instances updated.", modified_count)
    return modified_count
