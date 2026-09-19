import json
import time
from pathlib import Path
from typing import Dict, Any, List, Optional, Callable, Sequence
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from openai import OpenAI

from pptx_jahat.config import Config, DATA_DIR, OUTPUT_DIR
from pptx_jahat.tools.json_parser import safe_json_loads
from pptx_jahat.tools.image_gen import generate_image
from pptx_jahat.tools.preview import render_pptx_file_previews, image_to_base64_jpeg
from pptx_jahat.tools.pptx_builder import (
    _safe_update_text_frame,
    _replace_image_in_shape,
    _remove_shapes,
    clone_slide_across_presentations,
    run_slidecheck_qa,
    verify_and_auto_heal_pptx,
    _set_paragraph_rtl_and_fonts,
    _set_run_rtl_and_fonts,
    _is_rtl_text
)


# -----------------------------------------------------------------------------
# 1. INSPECTION: Extract Structured Information from PPTX for AI Editing
# -----------------------------------------------------------------------------
def inspect_pptx_for_editing(pptx_path: Path | str) -> Dict[str, Any]:
    """
    Extracts high-fidelity structural representation of a PowerPoint presentation
    including slide titles, shapes, text frames, tables, and speaker notes
    so the AI Agent can reason about requested edits with precision.
    """
    p = Path(pptx_path)
    if not p.exists():
        raise FileNotFoundError(f"Presentation not found at {p}")

    prs = Presentation(str(p))
    slides_info: List[Dict[str, Any]] = []

    for s_idx, slide in enumerate(prs.slides):
        slide_title = ""
        shapes_info: List[Dict[str, Any]] = []

        # Extract speaker notes if available
        notes_text = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass

        for sh_idx, shape_item in enumerate(slide.shapes):
            shape: Any = shape_item
            sh_type_name = "UNKNOWN"
            try:
                sh_type_name = str(shape.shape_type).split(".")[-1]
            except Exception:
                pass

            is_ph = getattr(shape, "is_placeholder", False)
            ph_type_str = ""
            if is_ph:
                try:
                    ph_type_str = str(shape.placeholder_format.type).split(".")[-1]
                except Exception:
                    pass

            sh_text = ""
            paragraphs: List[str] = []
            if getattr(shape, "has_text_frame", False):
                sh_text = shape.text_frame.text.strip()
                paragraphs = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
                if not slide_title and (sh_type_name == "TITLE" or "TITLE" in ph_type_str or (sh_idx == 0 and len(sh_text) < 100)):
                    slide_title = sh_text

            table_data = []
            if getattr(shape, "has_table", False):
                try:
                    for row in shape.table.rows:
                        r_vals = [cell.text.strip() for cell in row.cells]
                        table_data.append(r_vals)
                except Exception:
                    pass

            shapes_info.append({
                "shape_index": sh_idx,
                "name": shape.name,
                "type": sh_type_name,
                "is_placeholder": is_ph,
                "placeholder_type": ph_type_str,
                "text": sh_text,
                "paragraphs": paragraphs,
                "has_table": getattr(shape, "has_table", False),
                "table_data": table_data,
                "is_picture": shape.shape_type == MSO_SHAPE_TYPE.PICTURE or "PICTURE" in ph_type_str
            })

        slides_info.append({
            "slide_index": s_idx,
            "slide_number": s_idx + 1,
            "title": slide_title or f"Slide {s_idx + 1}",
            "notes": notes_text,
            "shapes_count": len(shapes_info),
            "shapes": shapes_info
        })

    return {
        "file_path": str(p),
        "filename": p.name,
        "slide_count": len(prs.slides),
        "slides": slides_info
    }


# -----------------------------------------------------------------------------
# 2. STEP 2 HUMAN TOUCH: Refine Extracted Content from Inputs with AI
# -----------------------------------------------------------------------------
def refine_extracted_content_with_ai(
    current_content: Dict[str, Any],
    user_prompt: str,
    log_cb: Optional[Callable[[str], None]] = None,
    timeout: Optional[float] = None
) -> Dict[str, Any]:
    """
    Reruns/refines the extracted slide content (from multi-modal inputs)
    based on the user's custom feedback or prompt.
    Allows user to modify slide counts, titles, bullet density, or focus topics.
    """
    def log(msg: str):
        if log_cb:
            log_cb(msg)

    effective_timeout = float(timeout or Config.LLM_TIMEOUT)
    log(f"[*] Human Touch Refinement: Refining extracted content with 9Router AI (timeout={effective_timeout}s)...")

    client = OpenAI(
        api_key=Config.NINEROUTER_KEY or "dummy_key",
        base_url=Config.get_openai_base_url(),
        timeout=effective_timeout
    )

    doc_title = current_content.get("document_title", "Presentation")
    sections = current_content.get("sections", [])
    raw_text = current_content.get("raw_text", "")
    sources_summary = current_content.get("sources_summary", [])

    system_prompt = (
        "You are PrismPresenter Slide Content Extraction Refinement Agent. "
        "You receive existing extracted presentation slide sections and a user's custom modification prompt. "
        "Your task is to REVISE, CONDENSE, EXPAND, REORGANIZE, or EDIT the slide sections strictly according to the user prompt.\n\n"
        "Guidelines:\n"
        "- Maintain coherent slide sections with meaningful titles, paragraphs, and punchy bullet points.\n"
        "- Respect user requests for language translations, topic focus, or slide count adjustments.\n"
        "- Return ONLY valid JSON adhering strictly to the schema."
    )

    user_message = f"""
Existing Document Title: {doc_title}
User Refinement Prompt:
{user_prompt}

Current Extracted Sections:
```json
{json.dumps({"document_title": doc_title, "sections": sections[:15]}, ensure_ascii=False, indent=2)}
```

Raw Source Notes Snippet:
{raw_text[:2000] if raw_text else "(None)"}

Output ONLY valid JSON adhering to this exact schema:
{{
  "document_title": "{doc_title}",
  "total_sections": 4,
  "sections": [
    {{
      "title": "Section Title",
      "paragraphs": ["Paragraph 1..."],
      "bullets": ["Bullet 1", "Bullet 2"]
    }}
  ]
}}
"""

    try:
        response = client.chat.completions.create(
            model=Config.NINEROUTER_CHAT_MODEL,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_message}
            ],
            temperature=0.2
        )
        content = response.choices[0].message.content or "{}"
        parsed = safe_json_loads(content)

        if parsed and "sections" in parsed and isinstance(parsed["sections"], list) and len(parsed["sections"]) > 0:
            log(f"[✓] Human Touch: AI successfully refined extracted content into {len(parsed['sections'])} sections.")
            return {
                "document_title": parsed.get("document_title", doc_title),
                "total_sections": len(parsed["sections"]),
                "sections": parsed["sections"],
                "raw_paragraphs": [p for s in parsed["sections"] for p in s.get("paragraphs", [])],
                "sources_summary": sources_summary,
                "raw_text": raw_text
            }
    except Exception as e:
        log(f"[!] Human Touch Refinement notice: {e}. Retaining current content.")

    return current_content


# -----------------------------------------------------------------------------
# 3. STEP 2.5 HUMAN TOUCH: Refine Restructured Slides with AI
# -----------------------------------------------------------------------------
def refine_restructured_slides_with_ai(
    current_restructured: Dict[str, Any],
    user_prompt: str,
    structure_blueprint: Optional[str] = None,
    log_cb: Optional[Callable[[str], None]] = None,
    timeout: Optional[float] = None
) -> Dict[str, Any]:
    """
    Reruns/refines the restructured storyboard slides (quadrants, animation steps,
    mathematical elements, teacher callouts, rewritten bullets, notes)
    based on the user's custom feedback or prompt.
    """
    def log(msg: str):
        if log_cb:
            log_cb(msg)

    effective_timeout = float(timeout or Config.LLM_TIMEOUT)
    log(f"[*] Human Touch Refinement: Refining storyboard slides with 9Router AI (timeout={effective_timeout}s)...")

    client = OpenAI(
        api_key=Config.NINEROUTER_KEY or "dummy_key",
        base_url=Config.get_openai_base_url(),
        timeout=effective_timeout
    )

    doc_title = current_restructured.get("document_title", "Presentation")
    slides = current_restructured.get("slides", [])

    system_prompt = (
        "You are PrismPresenter Storyboard Restructure Refinement Agent. "
        "You receive existing restructured presentation storyboard slides and a user's custom refinement prompt. "
        "Your task is to refine, update, rewrite, or rearrange the slides according to the user prompt "
        "while strictly preserving the storyboard specification schema (quadrants, circled animation steps ①②③, "
        "mathematical formulas, teacher callouts, rewritten bullets, speaker notes).\n\n"
        "Return ONLY valid JSON adhering to the schema."
    )

    schema_context = f"\nStoryboard Schema Specification (`structure.md`):\n{structure_blueprint[:1500]}\n" if structure_blueprint else ""

    user_message = f"""
{schema_context}
User Refinement Prompt:
{user_prompt}

Current Restructured Storyboard:
```json
{json.dumps(current_restructured, ensure_ascii=False, indent=2)}
```

Output ONLY valid JSON adhering to this exact schema:
{{
  "document_title": "{doc_title}",
  "total_slides": 4,
  "slides": [
    {{
      "slide_number": 1,
      "title": "Slide Title",
      "quadrant": "TR",
      "slide_type": "Theory / Definition",
      "core_concept": "Brief 1-line concept",
      "animation_steps": [
        {{"step": 1, "indicator": "①", "text": "Step 1 text"}}
      ],
      "mathematical_elements": {{
        "definitions": "Definition rule",
        "formulas": "$E=mc^2$",
        "sets": ""
      }},
      "visual_annotations": {{
        "diagrams": "",
        "eliminations": "",
        "confirmations": "",
        "key_invariants": "",
        "teacher_callouts": "نکته کلیدی..."
      }},
      "rewritten_bullets": [
        "① First point",
        "② Second point"
      ],
      "speaker_notes": "Presenter notes..."
    }}
  ]
}}
"""

    try:
        response = client.chat.completions.create(
            model=Config.NINEROUTER_CHAT_MODEL,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_message}
            ],
            temperature=0.2
        )
        content = response.choices[0].message.content or "{}"
        parsed = safe_json_loads(content)

        if parsed and "slides" in parsed and isinstance(parsed["slides"], list) and len(parsed["slides"]) > 0:
            log(f"[✓] Human Touch: AI successfully refined restructured storyboard ({len(parsed['slides'])} slides).")
            return parsed
    except Exception as e:
        log(f"[!] Human Touch Storyboard Refinement notice: {e}. Retaining current slides.")

    return current_restructured


# -----------------------------------------------------------------------------
# 4. AFTER DONE HUMAN TOUCH: Edit PPTX with AI Workflow
# -----------------------------------------------------------------------------
def edit_pptx_with_ai(
    pptx_path: Path | str,
    user_prompt: str,
    template_name: Optional[str] = None,
    log_cb: Optional[Callable[[str], None]] = None,
    timeout: Optional[float] = None
) -> Dict[str, Any]:
    """
    Allows user to view the generated presentation and reply with a custom prompt
    to rerun and edit the presentation in the AI workflow.
    Inspects existing slides, calls 9Router AI to devise exact modification actions,
    applies the edits to the PPTX presentation, verifies with SlideCheck QA,
    and returns the updated preview URLs and modification summary.
    """
    def log(msg: str):
        if log_cb:
            log_cb(msg)

    p = Path(pptx_path)
    if not p.exists():
        raise FileNotFoundError(f"PPTX presentation not found at: {p}")

    effective_timeout = float(timeout or Config.LLM_TIMEOUT)
    log(f"[*] Human Touch AI Editor: Inspecting presentation {p.name}...")

    # Step 1: Inspect presentation
    deck_info = inspect_pptx_for_editing(p)
    slide_count = deck_info["slide_count"]
    log(f"[✓] Deck inspected: {slide_count} slide(s) found.")

    # Step 2: Formulate AI Editing Plan with 9Router
    client = OpenAI(
        api_key=Config.NINEROUTER_KEY or "dummy_key",
        base_url=Config.get_openai_base_url(),
        timeout=effective_timeout
    )

    system_prompt = (
        "You are PrismPresenter Autonomous Presentation Editing Agent. "
        "You receive the exact structural inspection of an existing PowerPoint (.pptx) deck "
        "(all slides, shape indices, titles, texts, tables, and notes) "
        "along with the user's custom edit prompt.\n\n"
        "Your task is to formulate a precise list of actions to modify the presentation.\n"
        "Supported actions:\n"
        "1. 'update_text': Change the text of a specific shape on a slide.\n"
        "   Required fields: 'slide_index' (0-based), 'shape_index' (0-based), 'new_text'.\n"
        "2. 'update_table': Update table cell data.\n"
        "   Required fields: 'slide_index', 'shape_index', 'table_data' (2D array of strings).\n"
        "3. 'update_notes': Update speaker notes.\n"
        "   Required fields: 'slide_index', 'notes'.\n"
        "4. 'delete_slide': Remove an unwanted slide.\n"
        "   Required fields: 'slide_index'.\n"
        "5. 'reorder_slides': Reorder all slides.\n"
        "   Required fields: 'new_order' (array of 0-based slide indices, e.g. [1, 0, 2]).\n"
        "6. 'generate_image': Generate new AI image for a picture shape.\n"
        "   Required fields: 'slide_index', 'shape_index', 'prompt'.\n"
        "7. 'add_slide': Add a new slide copied from an existing slide or template with new text.\n"
        "   Required fields: 'copy_from_slide_index' (or 0), 'title', 'content_text', 'speaker_notes'.\n\n"
        "Return ONLY valid JSON with 'summary_of_changes' and 'actions' array."
    )

    deck_summary_simplified = []
    for s in deck_info["slides"]:
        shapes_summary = []
        for sh in s["shapes"]:
            if sh["text"] or sh["has_table"] or sh["is_picture"]:
                shapes_summary.append({
                    "shape_index": sh["shape_index"],
                    "name": sh["name"],
                    "type": sh["type"],
                    "text": sh["text"][:300] if sh["text"] else "",
                    "has_table": sh["has_table"],
                    "table_data": sh["table_data"],
                    "is_picture": sh["is_picture"]
                })
        deck_summary_simplified.append({
            "slide_index": s["slide_index"],
            "slide_number": s["slide_number"],
            "title": s["title"],
            "notes": s["notes"],
            "shapes": shapes_summary
        })

    user_message = f"""
Presentation File: {p.name}
Total Slides: {slide_count}

Deck Content:
```json
{json.dumps(deck_summary_simplified, ensure_ascii=False, indent=2)}
```

User Edit Prompt:
{user_prompt}

Output ONLY valid JSON adhering strictly to:
{{
  "summary_of_changes": "Clear 1-2 sentence description of what was updated.",
  "actions": [
    {{
      "action": "update_text",
      "slide_index": 0,
      "shape_index": 1,
      "new_text": "Updated text..."
    }}
  ]
}}
"""

    log(f"[*] Dispatching edit prompt to 9Router AI Model '{Config.NINEROUTER_CHAT_MODEL}'...")
    try:
        response = client.chat.completions.create(
            model=Config.NINEROUTER_CHAT_MODEL,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_message}
            ],
            temperature=0.2
        )
        content = response.choices[0].message.content or "{}"
        plan = safe_json_loads(content)
    except Exception as api_err:
        log(f"[!] AI Edit Reasoning error: {api_err}")
        return {
            "success": False,
            "error": f"AI Editing error: {str(api_err)}",
            "file_path": str(p),
            "summary": "AI call failed"
        }

    summary = plan.get("summary_of_changes", "Edits applied by AI.")
    actions = plan.get("actions", [])
    log(f"[✓] AI formulated {len(actions)} edit action(s): {summary}")

    # Step 3: Apply the actions to python-pptx presentation
    prs = Presentation(str(p))
    modified = False

    # A. First handle reordering or additions if present
    for act in actions:
        atype = act.get("action")
        if atype == "reorder_slides":
            new_order = act.get("new_order", [])
            if isinstance(new_order, list) and len(new_order) == len(prs.slides):
                try:
                    old_sld_ids = list(prs.slides._sldIdLst)
                    prs.slides._sldIdLst.clear()
                    for idx in new_order:
                        prs.slides._sldIdLst.append(old_sld_ids[idx])
                    modified = True
                    log(f"[✓] Reordered slides to: {new_order}")
                except Exception as ex:
                    log(f"[!] Reorder slides notice: {ex}")

        elif atype == "add_slide":
            copy_idx = act.get("copy_from_slide_index", 0)
            if copy_idx >= len(prs.slides):
                copy_idx = 0
            try:
                new_slide = clone_slide_across_presentations(prs, prs, copy_idx)
                title_text = act.get("title", "New Slide")
                content_text = act.get("content_text", "")
                notes_text = act.get("speaker_notes", "")

                text_shapes = [sh for sh in new_slide.shapes if sh.has_text_frame]
                if text_shapes:
                    _safe_update_text_frame(text_shapes[0].text_frame, title_text)
                    if len(text_shapes) > 1 and content_text:
                        _safe_update_text_frame(text_shapes[1].text_frame, content_text)

                if notes_text:
                    try:
                        new_slide.notes_slide.notes_text_frame.text = str(notes_text)
                    except Exception:
                        pass

                modified = True
                log(f"[✓] Added new slide: '{title_text}'")
            except Exception as ex:
                log(f"[!] Add slide notice: {ex}")

    # B. Next handle slide deletions (descending order)
    delete_indices = sorted([act.get("slide_index") for act in actions if act.get("action") == "delete_slide" and act.get("slide_index") is not None], reverse=True)
    for d_idx in delete_indices:
        if 0 <= d_idx < len(prs.slides):
            try:
                rId = prs.slides._sldIdLst[d_idx].rId
                prs.part.drop_rel(rId)
                prs.slides._sldIdLst.remove(prs.slides._sldIdLst[d_idx])
                modified = True
                log(f"[✓] Deleted slide index {d_idx}")
            except Exception as ex:
                log(f"[!] Delete slide notice: {ex}")

    # C. Handle in-slide updates: text, tables, notes, images
    for act in actions:
        atype = act.get("action")
        s_idx = act.get("slide_index")

        if s_idx is None or s_idx < 0 or s_idx >= len(prs.slides):
            continue

        slide = prs.slides[s_idx]

        if atype == "update_text":
            sh_idx = act.get("shape_index")
            new_text = act.get("new_text")
            if sh_idx is not None and 0 <= sh_idx < len(slide.shapes) and new_text is not None:
                shape: Any = slide.shapes[sh_idx]
                if getattr(shape, "has_text_frame", False):
                    _safe_update_text_frame(
                        shape.text_frame,
                        str(new_text),
                        is_rtl=None,
                        max_box_width_emu=getattr(shape, "width", None),
                        max_box_height_emu=getattr(shape, "height", None)
                    )
                    modified = True
                    log(f"[✓] Slide {s_idx + 1} shape #{sh_idx}: Updated text.")

        elif atype == "update_table":
            sh_idx = act.get("shape_index")
            tdata = act.get("table_data")
            if sh_idx is not None and 0 <= sh_idx < len(slide.shapes) and tdata:
                shape: Any = slide.shapes[sh_idx]
                if getattr(shape, "has_table", False):
                    for r_i, row in enumerate(tdata):
                        if r_i < len(shape.table.rows):
                            for c_i, cell_val in enumerate(row):
                                if c_i < len(shape.table.columns):
                                    cell = shape.table.cell(r_i, c_i)
                                    cell.text = str(cell_val)
                                    if cell.text_frame and cell.text_frame.paragraphs:
                                        p_elem = cell.text_frame.paragraphs[0]
                                        _set_paragraph_rtl_and_fonts(p_elem, is_rtl=_is_rtl_text(str(cell_val)))
                    modified = True
                    log(f"[✓] Slide {s_idx + 1} shape #{sh_idx}: Updated table data.")

        elif atype == "update_notes":
            notes_str = act.get("notes")
            if notes_str is not None:
                try:
                    notes_slide: Any = getattr(slide, "notes_slide", None)
                    if notes_slide and hasattr(notes_slide, "notes_text_frame") and notes_slide.notes_text_frame:
                        notes_slide.notes_text_frame.text = str(notes_str)
                        modified = True
                        log(f"[✓] Slide {s_idx + 1}: Updated speaker notes.")
                except Exception as ex:
                    log(f"[!] Speaker notes update notice: {ex}")

        elif atype == "generate_image":
            sh_idx = act.get("shape_index")
            img_prompt = act.get("prompt")
            if sh_idx is not None and 0 <= sh_idx < len(slide.shapes) and img_prompt:
                shape: Any = slide.shapes[sh_idx]
                try:
                    log(f"[*] Generating AI image for slide {s_idx + 1} shape #{sh_idx}: '{img_prompt[:40]}...'")
                    img_file = generate_image(img_prompt)
                    if img_file and not img_file.startswith("Error"):
                        if _replace_image_in_shape(shape, img_file):
                            modified = True
                            log(f"[✓] Slide {s_idx + 1} shape #{sh_idx}: Replaced image.")
                except Exception as ex:
                    log(f"[!] Image generation notice: {ex}")

    if modified or actions:
        prs.save(str(p))
        log(f"[✓] Presentation saved successfully.")

    # Step 4: Verification & SlideCheck QA
    try:
        qa_report = run_slidecheck_qa(p, auto_heal=True)
        if qa_report.get("overflow_issues_healed", 0) > 0:
            log(f"[QA] Auto-healed {qa_report['overflow_issues_healed']} overflowing text box(es).")
    except Exception as ex:
        log(f"[QA Warning] SlideCheck notice: {ex}")

    try:
        verify_and_auto_heal_pptx(p)
    except Exception as ex:
        log(f"[Verification Warning] Notice: {ex}")

    # Step 5: Render new previews
    previews = []
    engine_name = "Renderer"
    try:
        imgs, engine_name = render_pptx_file_previews(p, target_width_px=800, return_engine_info=True)
        img_list = imgs if isinstance(imgs, list) else [imgs]
        for idx, img in enumerate(img_list):
            previews.append({
                "slide_index": idx,
                "data_url": image_to_base64_jpeg(img, quality=85)
            })
    except BaseException as ex:
        log(f"[Preview Warning] Preview render notice: {ex}")

    return {
        "success": True,
        "file_path": str(p),
        "filename": p.name,
        "summary": summary,
        "actions_applied": len(actions),
        "slide_count": len(prs.slides),
        "previews": previews,
        "engine_name": engine_name
    }
