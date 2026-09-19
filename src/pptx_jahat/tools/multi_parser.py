import os
import io
import json
import base64
from pathlib import Path
from typing import Dict, Any, List, Optional, Callable, Sequence
from pptx import Presentation
from openai import OpenAI

from pptx_jahat.config import Config
from pptx_jahat.tools.docx_parser import parse_docx
from pptx_jahat.tools.json_parser import safe_json_loads

SUPPORTED_DOC_EXTS = {".docx", ".doc"}
SUPPORTED_PPT_EXTS = {".pptx"}
SUPPORTED_TXT_EXTS = {".txt", ".md", ".json", ".csv", ".tsv", ".yaml", ".yml"}
SUPPORTED_IMG_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}
SUPPORTED_AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac"}

ALL_SUPPORTED_EXTS = (
    SUPPORTED_DOC_EXTS
    | SUPPORTED_PPT_EXTS
    | SUPPORTED_TXT_EXTS
    | SUPPORTED_IMG_EXTS
    | SUPPORTED_AUDIO_EXTS
)

def is_supported_file(filename: str) -> bool:
    ext = Path(filename).suffix.lower()
    return ext in ALL_SUPPORTED_EXTS

def get_file_type_category(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in SUPPORTED_DOC_EXTS:
        return "Word Document"
    elif ext in SUPPORTED_PPT_EXTS:
        return "PowerPoint Presentation"
    elif ext in SUPPORTED_TXT_EXTS:
        return "Text / Markdown Document"
    elif ext in SUPPORTED_IMG_EXTS:
        return "Image / Visual Diagram"
    elif ext in SUPPORTED_AUDIO_EXTS:
        return "Audio Recording / Speech"
    return "Generic Document"

def parse_pptx_content(file_path: Path | str) -> Dict[str, Any]:
    """
    Extracts structured content, slide text, tables, and notes from an existing PPTX file.
    """
    path = Path(file_path)
    prs = Presentation(str(path))
    
    doc_title = path.stem.replace("_", " ")
    sections = []
    raw_paragraphs = []
    
    for idx, slide in enumerate(prs.slides):
        slide_title = f"Slide {idx + 1}"
        slide_paras = []
        slide_bullets = []
        slide_tables = []
        
        # Check title
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            t = slide.shapes.title.text.strip()
            if t:
                slide_title = t
                if idx == 0:
                    doc_title = t
                    
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_text_frame", False) and getattr(shape, "text_frame", None):
                tf = getattr(shape, "text_frame")
                txt = tf.text.strip()
                if txt:
                    lines = [line.strip() for line in txt.split("\n") if line.strip()]
                    for l in lines:
                        if l.startswith("- ") or l.startswith("• ") or l.startswith("* "):
                            slide_bullets.append(l.lstrip("- •*").strip())
                        else:
                            slide_paras.append(l)
                        raw_paragraphs.append(l)
            elif getattr(shape, "has_table", False) and getattr(shape, "table", None):
                tbl = getattr(shape, "table")
                tdata = []
                for row in tbl.rows:
                    tdata.append([cell.text.strip() for cell in row.cells])
                if tdata:
                    slide_tables.append(tdata)

        # Speaker notes
        notes_text = ""
        try:
            if slide.notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass

        if notes_text:
            slide_paras.append(f"[Presenter Note: {notes_text}]")

        sections.append({
            "title": slide_title,
            "level": 1 if idx == 0 else 2,
            "paragraphs": slide_paras,
            "bullets": slide_bullets,
            "tables": slide_tables,
            "speaker_notes": notes_text
        })
        
    return {
        "document_title": doc_title,
        "sections": sections,
        "raw_paragraphs": raw_paragraphs
    }

def parse_text_content(file_path: Path | str) -> Dict[str, Any]:
    """
    Parses plain text or markdown files into structured sections.
    """
    path = Path(file_path)
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()

    lines = text.split("\n")
    doc_title = path.stem.replace("_", " ")
    sections = []
    current_sec = {
        "title": doc_title,
        "level": 1,
        "paragraphs": [],
        "bullets": [],
        "tables": []
    }
    raw_paragraphs = []

    for line in lines:
        sline = line.strip()
        if not sline:
            continue
        raw_paragraphs.append(sline)
        if sline.startswith("# ") or sline.startswith("## ") or sline.startswith("### "):
            if current_sec["paragraphs"] or current_sec["bullets"]:
                sections.append(current_sec)
            level = sline.count("#", 0, 4)
            sec_title = sline.lstrip("#").strip()
            if level == 1 and current_sec["title"] == doc_title:
                doc_title = sec_title
            current_sec = {
                "title": sec_title,
                "level": level,
                "paragraphs": [],
                "bullets": [],
                "tables": []
            }
        elif sline.startswith("- ") or sline.startswith("• ") or sline.startswith("* "):
            current_sec["bullets"].append(sline.lstrip("- •*").strip())
        else:
            current_sec["paragraphs"].append(sline)

    if current_sec["paragraphs"] or current_sec["bullets"]:
        sections.append(current_sec)

    if not sections:
        sections.append({
            "title": doc_title,
            "level": 1,
            "paragraphs": [text[:500]],
            "bullets": [],
            "tables": []
        })

    return {
        "document_title": doc_title,
        "sections": sections,
        "raw_paragraphs": raw_paragraphs
    }

def _unpack_transcription(transcription: str, default_title: str) -> tuple[str, List[Dict[str, Any]], List[str]]:
    """
    Intelligently unpacks vision OCR transcription whether it was returned
    as structured JSON, Markdown, or plain text.
    """
    import re
    cleaned = transcription.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json|markdown|text)?\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned)
        cleaned = cleaned.strip()

    # 1. If valid JSON was returned by vision model
    if cleaned.startswith("{") and cleaned.endswith("}"):
        try:
            data = safe_json_loads(cleaned)
            title = str(data.get("main_topic") or data.get("title") or default_title)
            sections: List[Dict[str, Any]] = []
            raw_paras: List[str] = []

            raw_secs = data.get("sections", [])
            if isinstance(raw_secs, list) and raw_secs:
                for idx, s in enumerate(raw_secs):
                    if isinstance(s, dict):
                        stitle = str(s.get("title") or s.get("section_id") or f"Section {idx+1}")
                        sparas: List[str] = []
                        if s.get("problem_statement"):
                            sparas.append(str(s["problem_statement"]))
                        if s.get("options_analysis") and isinstance(s["options_analysis"], list):
                            for opt in s["options_analysis"]:
                                if isinstance(opt, dict):
                                    num = opt.get("number", "")
                                    fact = opt.get("factorization", "")
                                    otype = opt.get("type", "")
                                    sparas.append(f"• {num}: {fact} — {otype}".strip(" —"))
                                elif isinstance(opt, str):
                                    sparas.append(f"• {opt}")
                        for k, v in s.items():
                            if k not in ("title", "section_id", "problem_statement", "options_analysis"):
                                if isinstance(v, (str, int, float)):
                                    sparas.append(f"{k}: {v}")
                                elif isinstance(v, list):
                                    for it in v:
                                        sparas.append(f"• {it}")
                        sections.append({
                            "title": stitle,
                            "level": 2,
                            "paragraphs": sparas,
                            "bullets": [],
                            "tables": []
                        })
                        raw_paras.extend(sparas)
            else:
                sparas = []
                for k, v in data.items():
                    if isinstance(v, (str, int, float)):
                        sparas.append(f"{k}: {v}")
                    elif isinstance(v, list):
                        for it in v:
                            sparas.append(f"• {it}")
                sections.append({
                    "title": title,
                    "level": 1,
                    "paragraphs": sparas,
                    "bullets": [],
                    "tables": []
                })
                raw_paras.extend(sparas)

            if sections:
                return title, sections, raw_paras
        except Exception:
            pass

    # 2. If Markdown or plain text
    lines = cleaned.split("\n")
    title = default_title
    sections = []
    current_sec = {
        "title": default_title,
        "level": 1,
        "paragraphs": [],
        "bullets": [],
        "tables": []
    }
    raw_paras = []
    for line in lines:
        sline = line.strip()
        if not sline:
            continue
        raw_paras.append(sline)
        if sline.startswith("# ") or sline.startswith("## ") or sline.startswith("### "):
            if current_sec["paragraphs"] or current_sec["bullets"]:
                sections.append(current_sec)
            level = sline.count("#", 0, 4)
            sec_title = sline.lstrip("#").strip()
            if level == 1 and current_sec["title"] == default_title:
                title = sec_title
            current_sec = {
                "title": sec_title,
                "level": level,
                "paragraphs": [],
                "bullets": [],
                "tables": []
            }
        elif sline.startswith("- ") or sline.startswith("• ") or sline.startswith("* "):
            current_sec["bullets"].append(sline.lstrip("- •*").strip())
        else:
            current_sec["paragraphs"].append(sline)

    if current_sec["paragraphs"] or current_sec["bullets"]:
        sections.append(current_sec)

    if not sections:
        sections.append({
            "title": title,
            "level": 1,
            "paragraphs": [cleaned],
            "bullets": [],
            "tables": []
        })

    return title, sections, raw_paras

def parse_image_content(
    file_path: Path | str,
    log_cb: Optional[Callable[[str], None]] = None,
    timeout: Optional[float] = None,
    structure_blueprint: Optional[str] = None
) -> Dict[str, Any]:
    """
    Transcribes text, formulas, diagrams, and handwriting from an image file
    using 9Router Vision AI, optionally guided by a detection/storyboard structure.md schema.
    """
    def log(msg: str):
        if log_cb:
            log_cb(msg)

    path = Path(file_path)
    effective_timeout = float(timeout or Config.LLM_TIMEOUT)
    log(f"[*] Analyzing image content using Vision AI: {path.name} (timeout={effective_timeout}s)...")

    # Encode image to base64
    with open(path, "rb") as f:
        b64 = base64.b64encode(f.read()).decode("utf-8")
    
    ext = path.suffix.lower().lstrip(".")
    mime = "image/jpeg" if ext in ("jpg", "jpeg") else (f"image/{ext}" if ext else "image/png")
    data_uri = f"data:{mime};base64,{b64}"

    client = OpenAI(
        api_key=Config.NINEROUTER_KEY or "dummy_key",
        base_url=Config.get_openai_base_url(),
        timeout=effective_timeout
    )

    structure_hint = ""
    if structure_blueprint and structure_blueprint.strip():
        structure_hint = f"""
STRICT DETECTION & STORYBOARD SPECIFICATION SCHEMA (`structure.md`):
Follow this schema to detect and organize the slide elements in this image:
- Identify quadrant positions (e.g. Quadrant_TR: Top-Right, Quadrant_TL: Top-Left, Quadrant_BR: Bottom-Right, Quadrant_BL: Bottom-Left).
- Transcribe animation sequences strictly following circled numbers (①, ②, ③, ④...).
- Transcribe mathematical elements with formal LaTeX formatting (e.g., $91 = 7 \\times 13$).
- Detect visual/non-textual annotations: eliminations (crossed-out items), confirmations (checkmarked items), key boxed invariants, and teacher callouts (e.g., 'اگه دقت کنی!', 'نکته').
Schema reference:
{structure_blueprint[:1500]}
"""

    system_prompt = (
        "You are an expert Document OCR, Handwritten Math & Visual Content Transcriber. "
        "Transcribe and analyze all visible content in the provided image (such as handwritten paper notes, "
        "worksheet pages, 4-quadrant slide sheets, diagrams, formulas, or book pages). "
        "Extract:\n"
        "1. Main Title / Topic\n"
        "2. Key Concepts & Definitions\n"
        "3. Step-by-step procedures or formulas (use standard LaTeX formatting, e.g. $a \\times b = c$)\n"
        "4. Numbered step indicators (①, ②, ③...), eliminations, boxed items, and teacher callouts.\n"
        "Format as clean, structured Markdown text with clear headings (# and ##) and bullet points."
    )

    user_content = [
        {
            "type": "text",
            "text": f"Transcribe and extract the full content of this image into structured Markdown headings, formulas, and bullet points.\n{structure_hint}"
        },
        {
            "type": "image_url",
            "image_url": {"url": data_uri}
        }
    ]

    try:
        p_model = Config.get_agent_model("parser")
        p_think = Config.get_agent_think_level("parser")
        response = Config.safe_chat_completion(
            client,
            "parser",
            model=p_model,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_content}
            ],
            temperature=0.2
        )
        transcription = response.choices[0].message.content or ""
        log(f"[✓] Successfully transcribed visual content from {path.name} with '{p_model}' (thinking: {p_think}).")
    except Exception as ex:
        log(f"[!] Vision AI transcription warning: {ex}. Using basic image descriptor.")
        transcription = f"Visual reference graphic from {path.name}."

    doc_title, sections, raw_paragraphs = _unpack_transcription(transcription, default_title=f"Visual Source: {path.stem}")

    return {
        "document_title": doc_title,
        "sections": sections,
        "raw_paragraphs": raw_paragraphs
    }

def parse_audio_content(
    file_path: Path | str,
    log_cb: Optional[Callable[[str], None]] = None,
    timeout: Optional[float] = None
) -> Dict[str, Any]:
    """
    Transcribes speech from audio recordings using AI Whisper or speech endpoint.
    """
    def log(msg: str):
        if log_cb:
            log_cb(msg)

    path = Path(file_path)
    effective_timeout = float(timeout or Config.LLM_TIMEOUT)
    log(f"[*] Transcribing speech audio file: {path.name} (timeout={effective_timeout}s)...")

    client = OpenAI(
        api_key=Config.NINEROUTER_KEY or "dummy_key",
        base_url=Config.get_openai_base_url(),
        timeout=effective_timeout
    )

    transcript = ""
    try:
        # Try OpenAI-compatible audio transcription endpoint
        with open(path, "rb") as audio_file:
            res = client.audio.transcriptions.create(
                model="whisper-1",
                file=audio_file
            )
            transcript = getattr(res, "text", str(res))
            log(f"[✓] Audio transcription complete for {path.name}.")
    except Exception as ex:
        log(f"[!] Audio transcription warning ({ex}). Utilizing audio session metadata.")
        transcript = f"Audio lecture recording: {path.name}. Presentation audio notes."

    return {
        "document_title": f"Audio Lecture: {path.stem}",
        "sections": [
            {
                "title": f"Audio Recording: {path.name}",
                "level": 1,
                "paragraphs": [transcript],
                "bullets": [],
                "tables": []
            }
        ],
        "raw_paragraphs": [transcript]
    }

def parse_multiple_sources(
    file_paths: Sequence[str | Path],
    raw_text: Optional[str] = None,
    log_cb: Optional[Callable[[str], None]] = None,
    timeout: Optional[float] = None,
    structure_blueprint: Optional[str] = None
) -> Dict[str, Any]:
    """
    Unified multi-format and multi-file content extractor.
    Accepts any combination of:
    - Word (.docx, .doc)
    - PowerPoint (.pptx)
    - Text / Markdown (.txt, .md, .json)
    - Images (.png, .jpg, .webp) via Vision AI
    - Audio (.mp3, .wav, etc.) via Speech AI
    - Direct pasted raw text / notes
    Aggregates all parsed material into a unified document structure, optionally guided
    by a detection structure.md schema.
    """
    def log(msg: str):
        if log_cb:
            log_cb(msg)

    effective_timeout = float(timeout or Config.LLM_TIMEOUT)
    all_sections: List[Dict[str, Any]] = []
    all_raw_paragraphs: List[str] = []
    sources_summary: List[Dict[str, Any]] = []
    overall_title: Optional[str] = None

    valid_paths = [Path(p) for p in file_paths if Path(p).exists()]
    log(f"[*] Ingesting {len(valid_paths)} source file(s)...")

    for p in valid_paths:
        ext = p.suffix.lower()
        cat = get_file_type_category(p.name)
        sources_summary.append({
            "filename": p.name,
            "category": cat,
            "size_kb": round(p.stat().st_size / 1024, 1)
        })

        try:
            if ext in SUPPORTED_DOC_EXTS:
                parsed = parse_docx(p)
            elif ext in SUPPORTED_PPT_EXTS:
                parsed = parse_pptx_content(p)
            elif ext in SUPPORTED_TXT_EXTS:
                parsed = parse_text_content(p)
            elif ext in SUPPORTED_IMG_EXTS:
                parsed = parse_image_content(p, log_cb=log, timeout=effective_timeout, structure_blueprint=structure_blueprint)
            elif ext in SUPPORTED_AUDIO_EXTS:
                parsed = parse_audio_content(p, log_cb=log, timeout=effective_timeout)
            else:
                parsed = parse_text_content(p)

            if not overall_title and parsed.get("document_title"):
                overall_title = parsed["document_title"]

            for sec in parsed.get("sections", []):
                all_sections.append(sec)
            for rp in parsed.get("raw_paragraphs", []):
                all_raw_paragraphs.append(rp)

        except Exception as e:
            log(f"[!] Error parsing source file {p.name}: {e}")

    # Add optional pasted raw text
    if raw_text and raw_text.strip():
        sources_summary.append({
            "filename": "Direct Text Input",
            "category": "Raw Text / Notes",
            "size_kb": round(len(raw_text) / 1024, 1)
        })
        lines = [l.strip() for l in raw_text.strip().split("\n") if l.strip()]
        all_sections.append({
            "title": "Supplementary Notes",
            "level": 1,
            "paragraphs": lines,
            "bullets": [],
            "tables": []
        })
        all_raw_paragraphs.extend(lines)

    if not overall_title:
        if valid_paths:
            overall_title = valid_paths[0].stem.replace("_", " ")
        else:
            overall_title = "Presentation Document"

    if not all_sections:
        all_sections.append({
            "title": overall_title,
            "level": 1,
            "paragraphs": ["Presentation content synthesized from multi-modal inputs."],
            "bullets": [],
            "tables": []
        })

    return {
        "document_title": overall_title,
        "total_sections": len(all_sections),
        "raw_paragraphs": all_raw_paragraphs,
        "sections": all_sections,
        "sources_summary": sources_summary,
        "raw_text": "\n\n".join(all_raw_paragraphs[:50])
    }
