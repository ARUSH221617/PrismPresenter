import json
import pytest
import shutil
import tempfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx_jahat.config import DATA_DIR, Config
from pptx_jahat.tools.human_touch import (
    inspect_pptx_for_editing,
    refine_extracted_content_with_ai,
    refine_restructured_slides_with_ai,
    edit_pptx_with_ai
)
from pptx_jahat.web.app import create_app, JOBS, JOBS_LOCK
import threading

@pytest.fixture(autouse=True)
def configure_test_env():
    prev_pil = Config.PURE_PIL_ACTIVE
    Config.PURE_PIL_ACTIVE = True
    yield
    Config.PURE_PIL_ACTIVE = prev_pil

@pytest.fixture
def client():
    app = create_app()
    app.config["TESTING"] = True
    with app.test_client() as client:
        yield client

@pytest.fixture
def sample_pptx(tmp_path):
    """Creates a clean sample PPTX with 2 slides for testing edits."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Title + subtitle
    s1 = prs.slides.add_slide(blank_layout)
    tb1 = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.5))
    tb1.text_frame.text = "Original Slide 1 Title"
    tb2 = s1.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(3))
    tb2.text_frame.text = "Original bullet 1\nOriginal bullet 2"

    # Slide 2: Table
    s2 = prs.slides.add_slide(blank_layout)
    table_shape = s2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(3))
    table = table_shape.table
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Data 1"
    table.cell(1, 1).text = "Data 2"

    out_path = tmp_path / "sample_test.pptx"
    prs.save(str(out_path))
    return out_path

def test_inspect_pptx_for_editing(sample_pptx):
    info = inspect_pptx_for_editing(sample_pptx)
    assert info["slide_count"] == 2
    assert len(info["slides"]) == 2
    assert info["slides"][0]["title"] == "Original Slide 1 Title"
    assert info["slides"][0]["shapes_count"] == 2
    assert info["slides"][1]["shapes"][0]["has_table"] is True
    assert info["slides"][1]["shapes"][0]["table_data"][0][0] == "Header A"

def test_edit_pptx_with_ai_mocked(sample_pptx, monkeypatch):
    """Test AI edit flow with a controlled plan mock."""
    from unittest.mock import MagicMock
    mock_plan = {
        "summary_of_changes": "Updated Slide 1 title and modified table data.",
        "actions": [
            {
                "action": "update_text",
                "slide_index": 0,
                "shape_index": 0,
                "new_text": "Updated AI Title"
            },
            {
                "action": "update_table",
                "slide_index": 1,
                "shape_index": 0,
                "table_data": [["Col 1", "Col 2"], ["Val 1", "Val 2"]]
            },
            {
                "action": "update_notes",
                "slide_index": 0,
                "notes": "Speaker notes for slide 1"
            }
        ]
    }

    mock_client = MagicMock()
    mock_resp = MagicMock()
    mock_resp.choices = [MagicMock(message=MagicMock(content=json.dumps(mock_plan)))]
    mock_client.chat.completions.create.return_value = mock_resp
    monkeypatch.setattr("pptx_jahat.tools.human_touch.OpenAI", lambda **kwargs: mock_client)

    res = edit_pptx_with_ai(sample_pptx, "Change title and table")
    assert res["success"] is True
    assert res["actions_applied"] == 3

    # Inspect modified presentation
    re_inspected = inspect_pptx_for_editing(sample_pptx)
    assert re_inspected["slides"][0]["shapes"][0]["text"] == "Updated AI Title"
    assert re_inspected["slides"][0]["notes"] == "Speaker notes for slide 1"
    assert re_inspected["slides"][1]["shapes"][0]["table_data"][0][0] == "Col 1"

def test_api_extract_endpoint(client):
    res = client.post("/api/generator/extract", json={
        "raw_text": "Introduction\nThis is paragraph one.\n- Bullet point one\n- Bullet point two"
    })
    assert res.status_code == 200
    data = res.get_json()
    assert data["success"] is True
    assert "data" in data
    assert data["data"]["total_sections"] >= 1

def test_api_inspect_deck_endpoint(client, sample_pptx):
    res = client.post("/api/generator/inspect-deck", json={
        "file_path": str(sample_pptx)
    })
    assert res.status_code == 200
    data = res.get_json()
    assert data["success"] is True
    assert data["deck"]["slide_count"] == 2

def test_api_edit_pptx_endpoint(client, sample_pptx, monkeypatch):
    from unittest.mock import MagicMock
    mock_plan = {
        "summary_of_changes": "Updated Slide 1 text",
        "actions": [
            {
                "action": "update_text",
                "slide_index": 0,
                "shape_index": 0,
                "new_text": "New API Text"
            }
        ]
    }
    mock_client = MagicMock()
    mock_resp = MagicMock()
    mock_resp.choices = [MagicMock(message=MagicMock(content=json.dumps(mock_plan)))]
    mock_client.chat.completions.create.return_value = mock_resp
    monkeypatch.setattr("pptx_jahat.tools.human_touch.OpenAI", lambda **kwargs: mock_client)

    res = client.post("/api/generator/edit-pptx", json={
        "file_path": str(sample_pptx),
        "prompt": "Update text"
    })
    assert res.status_code == 200
    data = res.get_json()
    assert data["success"] is True
    assert data["actions_applied"] == 1

def test_api_human_action_signaling(client):
    job_id = "test_human_job_123"
    event = threading.Event()
    with JOBS_LOCK:
        JOBS[job_id] = {
            "id": job_id,
            "waiting_for_human": True,
            "human_step": "extract",
            "human_event": event,
            "step_data": {"document_title": "Test Title", "sections": []}
        }

    # Query human status
    status_res = client.get(f"/api/generator/human-status/{job_id}")
    assert status_res.status_code == 200
    assert status_res.get_json()["waiting_for_human"] is True

    # Send continue action
    res = client.post("/api/generator/human-action", json={
        "job_id": job_id,
        "action": "continue",
        "step": "extract",
        "data": {"document_title": "Approved Title", "sections": []}
    })
    assert res.status_code == 200
    assert res.get_json()["success"] is True
    assert event.is_set()

    with JOBS_LOCK:
        assert JOBS[job_id]["human_action"]["action"] == "continue"
        assert JOBS[job_id]["human_action"]["data"]["document_title"] == "Approved Title"
        del JOBS[job_id]
